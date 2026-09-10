"""S2 광역 탐색(로컬) + 시설명 추출 + 검수용 Excel / S3 시설명 재검색.

전량 수집된 notices_all.parquet 위에서만 동작하므로 API 호출이 없다.
- 시설 후보 단위 = (표2 분류, 검색어, 시설키, 수요기관). 같은 이름의 시설이 여러 지자체에 있어도 섞이지 않도록 수요기관을 키에 포함한다.
- discover 를 다시 실행해도 기존 facility_candidates.xlsx 의 검수 내용(노란 셀)과 시설ID는 (시설키, 수요기관) 기준으로 유지된다.
"""
from __future__ import annotations

import logging
import os
import re
from typing import Dict, List, Optional

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .classify import (classify_notice_kind, classify_trade, classify_work_type, extract_facility_name, load_keywords,
                       match_categories, normalize_facility_key, parse_amount, prefix_institution, strip_institution_prefix,
                       work_type_for)

log = logging.getLogger(__name__)

YELLOW = PatternFill("solid", fgColor="FFFF00")
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
REVIEW_COLS = ["검수_포함여부", "검수_시설명", "검수_별칭(;구분)", "검수_수요기관", "검수_메모"]
_EMPTY_STRS = {"", "nan", "none", "null", "nat", "<na>"}


def _s(series: pd.Series) -> pd.Series:
    """원문 값 → 문자열('nan'/'None'/공백은 공란). 한 번의 map 으로 처리(129만 행 규모에서 두 번 map 은 비용이 큼)."""
    def conv(v):
        if v is None or (isinstance(v, float) and v != v):
            return ""
        t = str(v)
        return "" if t.strip().lower() in _EMPTY_STRS else t
    return series.map(conv)


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and v != v) or str(v).strip().lower() in _EMPTY_STRS


def raw_columns_needed(cfg: dict) -> List[str]:
    """standardize 가 실제로 읽는 원문 컬럼(전량 parquet 는 145개 컬럼 × 백만 행이므로 필요한 것만 읽어 메모리를 줄인다)."""
    f = cfg["fields"]
    keys = ["bid_no", "bid_ord", "name", "kind", "re_notice", "reg_type", "notice_dt", "close_dt", "open_dt", "notice_inst",
            "demand_inst", "presmpt_price", "main_cnstty", "site_region", "detail_url", "ref_no", "prespec_no",
            "demand_inst_cd", "bdgt_amt", "vat", "govsply_amt", "contractor_govsply", "gov_govsply", "prev_bid_no", "chg_reason",
            "sucsfbid_rate", "std_notice_url"]
    cols = [f[k] for k in keys if isinstance(f.get(k), str) and f[k]]
    cols += [f"{f.get('sub_cnstty_prefix', '')}{i}" for i in range(1, int(f.get("sub_cnstty_max", 0) or 0) + 1)]
    cols += [f"{f.get('spt_url_prefix', '')}{i}" for i in range(1, int(f.get("spt_max", 0) or 0) + 1)]
    n = int(f["attach_max"])
    cols += [f"{f['attach_url_prefix']}{i}" for i in range(1, n + 1)] + [f"{f['attach_name_prefix']}{i}" for i in range(1, n + 1)]
    return list(dict.fromkeys(cols))


def standardize(df_raw: pd.DataFrame, cfg: dict, bsis: Optional[pd.DataFrame] = None,
                license_df: Optional[pd.DataFrame] = None, awards: Optional[pd.DataFrame] = None) -> pd.DataFrame:
    """API 원문 → 내부 표준 컬럼. 미존재 필드는 공란."""
    f = cfg["fields"]
    empty = pd.Series([""] * len(df_raw), index=df_raw.index, dtype=object)
    g = lambda k: _s(df_raw[f[k]]) if f[k] in df_raw.columns else empty
    df = pd.DataFrame({
        "공고번호": g("bid_no"),
        "공고차수": g("bid_ord"),
        "공고명": g("name"),
        "공고종류_원문": g("kind"),
        "재공고여부": g("re_notice"),
        "등록유형": g("reg_type"),
        "공고일시": g("notice_dt"),
        "입찰마감일시": g("close_dt"),
        "개찰일시": g("open_dt"),
        "공고기관": g("notice_inst"),
        "수요기관": g("demand_inst"),
        "추정가격": g("presmpt_price").map(parse_amount),
        "주공종명": g("main_cnstty"),
        "공사현장지역": g("site_region"),
        "상세URL": g("detail_url"),
        "참조번호": g("ref_no"),
        "사전규격번호": g("prespec_no"),
    })
    # 실제 응답에서 확인된 추가 필드(없으면 공란)
    opt = lambda k: _s(df_raw[f[k]]) if (k in f and f[k] in df_raw.columns) else empty
    df["수요기관코드"] = opt("demand_inst_cd")
    df["예산금액"] = opt("bdgt_amt").map(parse_amount)
    df["부가세_API"] = opt("vat").map(parse_amount)
    df["관급자재_API"] = opt("govsply_amt").map(parse_amount)
    df["도급자관급액_API"] = opt("contractor_govsply").map(parse_amount)
    df["관급자관급액_API"] = opt("gov_govsply").map(parse_amount)
    df["이전공고번호"] = opt("prev_bid_no")
    df["변경사유"] = opt("chg_reason")
    df["낙찰하한율"] = opt("sucsfbid_rate")
    df["표준공고문URL"] = opt("std_notice_url")
    subs = []
    for i in range(1, int(f.get("sub_cnstty_max", 0) or 0) + 1):
        c = f"{f.get('sub_cnstty_prefix', '')}{i}"
        if c in df_raw.columns:
            subs.append(_s(df_raw[c]))
    df["부공종명"] = [" / ".join(x for x in row if x) for row in zip(*subs)] if subs else ""
    for i in range(1, int(f.get("spt_max", 0) or 0) + 1):
        c = f"{f.get('spt_url_prefix', '')}{i}"
        df[f"현장설명서URL{i}"] = _s(df_raw[c]) if c in df_raw.columns else ""
    # 첨부파일
    for i in range(1, int(f["attach_max"]) + 1):
        u, nm = f"{f['attach_url_prefix']}{i}", f"{f['attach_name_prefix']}{i}"
        df[f"첨부URL{i}"] = _s(df_raw[u]) if u in df_raw.columns else ""
        df[f"첨부파일명{i}"] = _s(df_raw[nm]) if nm in df_raw.columns else ""
    df["첨부파일수"] = sum((df[f"첨부URL{i}"].str.len() > 4).astype(int) for i in range(1, int(f["attach_max"]) + 1))
    df["공고차수"] = df["공고차수"].map(lambda s: s.zfill(3) if s.isdigit() else (s or "000"))   # 나라장터 차수는 '000' 3자리
    df["공고키"] = df["공고번호"] + "-" + df["공고차수"]

    # 기초금액 병합(공고번호+차수 기준, 없으면 공고번호)
    df["기초금액"] = None
    if bsis is not None and not bsis.empty and f["bid_no"] in bsis.columns:
        b = bsis.copy()
        if "_bsis_amount" not in b.columns:
            b["_bsis_amount"] = b[f["bsis_amount"]].map(parse_amount) if f["bsis_amount"] in b.columns else None
        ords = _s(b[f["bid_ord"]]) if f["bid_ord"] in b.columns else pd.Series(["000"] * len(b), index=b.index)
        b["_k"] = _s(b[f["bid_no"]]) + "-" + ords.map(lambda s: s.zfill(3) if s.isdigit() else (s or "000"))
        b = b[b["_bsis_amount"].notna()]
        m = b.drop_duplicates("_k", keep="last").set_index("_k")["_bsis_amount"]
        df["기초금액"] = df["공고키"].map(m)
        m2 = b.assign(_no=_s(b[f["bid_no"]])).drop_duplicates("_no", keep="last").set_index("_no")["_bsis_amount"]
        df["기초금액"] = df["기초금액"].fillna(df["공고번호"].map(m2))

    # 낙찰정보 병합(참고 컬럼: 공고번호+차수 → 없으면 공고번호). 같은 공고에 여러 행이면 개찰일시가 최신인 행
    for c in ("낙찰금액_API", "낙찰률_API", "낙찰자", "참가업체수", "낙찰개찰일시"):
        df[c] = None
    if awards is not None and not awards.empty and f["bid_no"] in awards.columns:
        from .collect import award_column
        w = awards.copy()
        amt_c, rate_c = award_column(cfg, w, "award_amt"), award_column(cfg, w, "award_rate")
        bid_c, cnt_c, dt_c = award_column(cfg, w, "award_bidder"), award_column(cfg, w, "award_prtcpt_cnt"), award_column(cfg, w, "award_open_dt")
        w["_amt"] = w["_award_amt"] if "_award_amt" in w.columns else (w[amt_c].map(parse_amount) if amt_c else None)
        w["_rate"] = w["_award_rate"] if "_award_rate" in w.columns else \
            (pd.to_numeric(_s(w[rate_c]).str.replace("%", "").str.replace(",", ""), errors="coerce") if rate_c else None)
        w["_bidder"] = _s(w[bid_c]) if bid_c else ""
        w["_cnt"] = pd.to_numeric(_s(w[cnt_c]), errors="coerce") if cnt_c else None
        w["_dt"] = _s(w[dt_c]) if dt_c else ""
        ords = _s(w[f["bid_ord"]]) if f["bid_ord"] in w.columns else pd.Series(["000"] * len(w), index=w.index)
        w["_no"] = _s(w[f["bid_no"]])
        w["_k"] = w["_no"] + "-" + ords.map(lambda s: s.zfill(3) if s.isdigit() else (s or "000"))
        w = w[w["_amt"].notna() | w["_rate"].notna()].sort_values("_dt")
        for key_col, src in (("_k", df["공고키"]), ("_no", df["공고번호"])):
            m = w.drop_duplicates(key_col, keep="last").set_index(key_col)
            for c, wc in (("낙찰금액_API", "_amt"), ("낙찰률_API", "_rate"), ("낙찰자", "_bidder"), ("참가업체수", "_cnt"), ("낙찰개찰일시", "_dt")):
                filled = src.map(m[wc])
                df[c] = df[c].where(df[c].notna(), filled)

    # 면허제한 업종(복수 → ' / ' 결합)
    df["면허제한업종"] = ""
    if license_df is not None and not license_df.empty and f["license_name"] in license_df.columns:
        lic = license_df.groupby("_bid_no")[f["license_name"]].apply(lambda s: " / ".join(sorted(set(map(str, s)))))
        df["면허제한업종"] = df["공고번호"].map(lic).fillna("")

    df["공고종류"] = [classify_notice_kind(k, r, t, n) for k, r, t, n in
                   zip(df["공고종류_원문"], df["재공고여부"], df["등록유형"], df["공고명"])]
    tr = [classify_trade(n, m, l) for n, m, l in zip(df["공고명"], df["주공종명"], df["면허제한업종"])]
    df["공종"] = [t[0] for t in tr]
    df["공종근거"] = [t[1] for t in tr]
    df["사업유형"] = df["공고명"].map(classify_work_type)
    return df


def _mode(s: pd.Series, default: str = "") -> str:
    vc = s[s.map(lambda v: not _blank(v))].value_counts()
    return str(vc.index[0]) if len(vc) else default


_PROJECT_TYPE_PRIORITY = ("신축", "증축·리모델링", "증축", "리모델링")


def _facility_work_type(s: pd.Series) -> str:
    """시설의 사업유형: 공고 하나라도 신축·증축·리모델링이면 그 유형(우선순위 순). 유지보수 공고가 아무리 많아도 본공사 1건을 묻지 않는다.
    프로젝트 유형이 없으면 미분류/유지보수 중 최빈값."""
    present = set(s.dropna().astype(str))
    for t in _PROJECT_TYPE_PRIORITY:
        if t in present:
            return t
    return _mode(s, "미분류")


def discover_candidates(std: pd.DataFrame, previous_review: Optional[pd.DataFrame] = None, ids_only: bool = False) -> pd.DataFrame:
    """공고명에 표2 검색어가 포함된 공고 → 시설명 후보 + 분류 (검수용).
    previous_review: 이전 facility_candidates.xlsx 내용(있으면 시설ID·검수 컬럼을 (시설키, 수요기관) 기준으로 이어받음).
    ids_only: 시설ID만 이어받고 검수 컬럼은 모두 새 기본값으로(`discover --fresh`)."""
    kw = load_keywords()
    rows = []
    # 공고명만 순회하고(iterrows 는 행마다 Series 를 만들어 백만 행 규모에서 느림) 매칭된 행만 전체 컬럼을 읽는다
    for i, name in enumerate(std["공고명"].tolist()):
        matches = match_categories(name, kw)
        if not matches:
            continue
        base = std.iloc[i].to_dict()
        for major, minor, w in matches:
            fac = prefix_institution(extract_facility_name(name, w), w, str(base.get("수요기관") or ""))   # 단지명 없으면 지자체명 접두
            rows.append({**base, "표2_대분류": major, "표2_중분류": minor, "검색어": w,
                         "시설명_후보": fac, "시설키": normalize_facility_key(fac),
                         "사업유형": work_type_for(name, w, fac)})   # 시설 자체가 주차장 등이면 부속어에서 제외
    cand = pd.DataFrame(rows)
    if cand.empty:
        return cand
    cand["수요기관"] = _s(cand["수요기관"])
    # 시설 단위 요약 (검수 대상은 공고가 아니라 시설). 같은 이름이라도 수요기관이 다르면 다른 시설로 본다.
    agg = (cand.groupby(["표2_대분류", "표2_중분류", "검색어", "시설키", "수요기관"], as_index=False)
           .agg(시설명_후보=("시설명_후보", lambda s: _mode(s)),
                공고건수=("공고번호", "nunique"),
                최초공고일=("공고일시", "min"), 최종공고일=("공고일시", "max"),
                사업유형=("사업유형", _facility_work_type),
                사업유형_분포=("사업유형", lambda s: ", ".join(f"{k} {v}" for k, v in s.value_counts().items())),
                공종목록=("공종", lambda s: "/".join(sorted(set(s)))),
                추정가격_합계=("추정가격", lambda s: pd.to_numeric(s, errors="coerce").fillna(0).sum()),
                공고명_예시=("공고명", "first")))
    agg = agg[["표2_대분류", "표2_중분류", "검색어", "시설키", "시설명_후보", "수요기관", "공고건수", "최초공고일", "최종공고일",
               "사업유형", "사업유형_분포", "공종목록", "추정가격_합계", "공고명_예시"]]

    agg["검수_포함여부"] = agg["사업유형"].map(default_include)
    agg["검수_시설명"] = agg["시설명_후보"]
    agg["검수_별칭(;구분)"] = ""
    agg["검수_수요기관"] = agg["수요기관"]
    agg["검수_메모"] = ""
    agg = _carry_over_review(agg, previous_review, ids_only=ids_only)
    cols = ["시설ID"] + [c for c in agg.columns if c != "시설ID"]
    return agg[cols].reset_index(drop=True)


def default_include(work_type: str) -> str:
    """사업유형 → 검수_포함여부 기본값."""
    if work_type in ("신축", "증축", "리모델링", "증축·리모델링"):
        return "포함"
    return "제외" if work_type == "유지보수" else "검토필요"


def _user_edited(prev_row, col: str) -> bool:
    """이전 파일의 검수 값이 사용자가 고친 것인지(기본값 그대로면 False).
    기본값이었다면 새 규칙으로 다시 계산한 값을 쓰고, 고친 값만 이어받는다."""
    v = prev_row.get(col)
    if _blank(v):
        return False
    v = str(v).strip()
    if col == "검수_포함여부" and not _blank(prev_row.get("사업유형")):
        return v != default_include(str(prev_row.get("사업유형")).strip())
    if col == "검수_시설명" and not _blank(prev_row.get("시설명_후보")):
        return v != str(prev_row.get("시설명_후보")).strip()
    if col == "검수_수요기관" and not _blank(prev_row.get("수요기관")):
        return v != str(prev_row.get("수요기관")).strip()
    return True          # 별칭·메모(기본값 공란)와, 비교 기준 컬럼이 없는 옛 형식 파일


def _carry_over_review(agg: pd.DataFrame, prev: Optional[pd.DataFrame], ids_only: bool = False) -> pd.DataFrame:
    """이전 검수 파일의 시설ID와 **사용자가 고친** 검수_* 값을 (시설키, 수요기관) 기준으로 이어받고, 새 시설에는 이어지는 번호를 부여.
    기본값 그대로였던 칸(예: 사업유형 최빈값에 따른 '제외')은 이어받지 않고 새 규칙의 기본값으로 다시 채운다."""
    agg = agg.copy()
    agg["시설ID"] = ""
    if prev is not None and not prev.empty and "시설키" in prev.columns and "시설ID" in prev.columns:
        p = prev.copy()
        p["수요기관"] = _s(p["수요기관"]) if "수요기관" in p.columns else ""
        p = p.drop_duplicates(["시설키", "수요기관"], keep="first").set_index(["시설키", "수요기관"])
        kept, edited = 0, 0
        for i, r in agg.iterrows():
            k = (r["시설키"], r["수요기관"])
            if k in p.index:
                old = p.loc[k]
                agg.at[i, "시설ID"] = str(old["시설ID"])
                any_edit = False
                for c in ([] if ids_only else REVIEW_COLS):
                    if c in p.columns and _user_edited(old, c):
                        agg.at[i, c] = str(old[c]).strip()
                        any_edit = True
                kept += 1
                edited += int(any_edit)
        log.info("이전 검수 내용 이어받음: %d개 시설(그중 사용자가 고친 값이 있는 시설 %d개, 나머지는 새 기본값)", kept, edited)
    used = set(agg.loc[agg["시설ID"] != "", "시설ID"])
    nums = [int(re.sub(r"\D", "", x) or 0) for x in used]
    nxt = (max(nums) if nums else 0) + 1
    for i in agg.index[agg["시설ID"] == ""]:
        agg.at[i, "시설ID"] = f"F{str(nxt).zfill(4)}"
        nxt += 1
    return agg


def write_review_workbook(agg: pd.DataFrame, path: str) -> None:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "시설후보_검수"
    cols = list(agg.columns)
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        ws.cell(row=1, column=c).font = BOLD
        ws.cell(row=1, column=c).fill = HEADER_FILL
    for row in agg.itertuples(index=False):
        ws.append([None if _blank(v) and not isinstance(v, (int, float)) else (v.item() if hasattr(v, "item") else v) for v in row])
    edit_cols = [i + 1 for i, c in enumerate(cols) if c.startswith("검수_")]
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(cols) + 1):
            ws.cell(row=r, column=c).font = FONT
        for c in edit_cols:
            ws.cell(row=r, column=c).fill = YELLOW
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(45, max(12, len(str(c)) + 4))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    legend = wb.create_sheet("범례")
    for line in ["노란색 셀만 수정하세요. (다른 열은 수정해도 반영되지 않습니다)",
                 "검수_포함여부: '포함' → DB 수집 대상 / '제외' → 제외 / '검토필요' → 판단 후 둘 중 하나로 바꾸기",
                 "기본값: 신축·증축·리모델링(증축·리모델링 포함)은 '포함', 유지보수(방수·도색·교체 등)는 '제외', 미분류는 '검토필요'",
                 "같은 시설에 신축과 리모델링이 모두 있으면 한 줄로 보입니다(사업유형_분포 참고). '포함'이면 두 사업이 각각 별도 프로젝트로 DB에 들어갑니다.",
                 "검수_시설명: 재검색에 사용할 정확한 시설명 (예: 수원시립미술관). 비워 두면 시설명_후보를 사용합니다.",
                 "검수_별칭(;구분): 재검색 보조어를 ';'로 구분 (예: 수원시립미술관;수원시 미술관;수원미술관)",
                 "검수_수요기관: 재검색 시 이 수요기관의 공고만 찾습니다(같은 이름의 시설이 다른 지자체에 있어도 섞이지 않게). 비우면 이름만으로 찾습니다.",
                 "사업유형이 '유지보수'로 잡혔지만 실제 리모델링·신축이면 '포함'으로 바꾸고 메모를 남기세요.",
                 "discover 를 다시 실행해도 이 파일의 검수 내용과 시설ID는 (시설키, 수요기관)이 같은 행에 그대로 이어집니다."]:
        legend.append([line])
    wb.save(path)
    log.info("검수용 워크북 저장: %s (%d개 시설)", path, len(agg))


def read_review_file(path: str) -> pd.DataFrame:
    """facility_candidates.xlsx 전체 행(필터 없음). 파일이 없으면 빈 DataFrame."""
    if not os.path.exists(path):
        return pd.DataFrame()
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["시설후보_검수"]
    rows = list(ws.values)
    wb.close()
    if not rows:
        return pd.DataFrame()
    return pd.DataFrame(rows[1:], columns=[str(c) for c in rows[0]])


def read_reviewed(path: str) -> pd.DataFrame:
    """검수 완료 행('포함')만. 검수_시설명이 비어 있으면 시설명_후보로 대체, 수요기관 컬럼이 없으면 빈 값."""
    df = read_review_file(path)
    if df.empty:
        raise SystemExit(f"검수 파일이 비어 있거나 없습니다: {path} (discover 를 먼저 실행)")
    for c in REVIEW_COLS + ["시설명_후보", "수요기관", "시설ID", "표2_대분류", "표2_중분류", "검색어"]:
        if c not in df.columns:
            df[c] = ""
    df = df[df["검수_포함여부"].map(lambda v: str(v).strip() == "포함")].copy()
    df["검수_시설명"] = [str(n).strip() if not _blank(n) else str(f).strip() for n, f in zip(df["검수_시설명"], df["시설명_후보"])]
    df["검수_별칭(;구분)"] = df["검수_별칭(;구분)"].map(lambda v: "" if _blank(v) else str(v))
    df["검수_수요기관"] = df["검수_수요기관"].map(lambda v: "" if _blank(v) else str(v).strip())
    df["시설ID"] = df["시설ID"].map(lambda v: str(v).strip())
    return df


def research_by_facility(std: pd.DataFrame, reviewed: pd.DataFrame, max_hits_warn: int = 300) -> pd.DataFrame:
    """검수된 시설명(+별칭)으로 전량 데이터를 재검색 → 모든 공종 공고 수집, 시설ID 부여.
    검수_수요기관이 있으면 수요기관이 일치(공백 무시, 부분 포함)하는 공고만 채택."""
    out = []
    std_nospace = std["공고명"].astype(str).str.replace(" ", "", regex=False)
    inst_nospace = std["수요기관"].astype(str).str.replace(" ", "", regex=False)
    inst_values = [v for v in inst_nospace.unique().tolist() if v]      # 수요기관 판정은 고유값(수천 개)에만 하고 isin 으로 확장
    for _, fac in reviewed.iterrows():
        names = [fac["검수_시설명"]] + [a.strip() for a in str(fac.get("검수_별칭(;구분)") or "").split(";") if a.strip()]
        inst_src = str(fac.get("검수_수요기관") or fac.get("수요기관") or "")
        # '구례군 국민임대아파트' 처럼 지자체명을 접두한 시설명은 공고명에 지자체명이 없으므로 접두어를 뗀 패턴도 함께 검색(수요기관 필터가 범위를 좁힌다)
        names += [strip_institution_prefix(n, inst_src) for n in list(names) if n and not _blank(n)]
        pats = list(dict.fromkeys(n.replace(" ", "") for n in names if n and not _blank(n)))
        if not pats:
            log.warning("시설 %s: 검수_시설명이 비어 있어 건너뜀", fac.get("시설ID"))
            continue
        if min(len(p) for p in pats) <= 2:
            log.warning("시설 %s '%s': 검색어가 너무 짧아(2자 이하) 엉뚱한 공고가 대량 포함될 수 있음", fac.get("시설ID"), "/".join(pats))
        mask = pd.Series(False, index=std.index)
        for p in pats:
            mask |= std_nospace.str.contains(re.escape(p), na=False)
        inst = str(fac.get("검수_수요기관") or "").replace(" ", "")
        if inst:
            mask &= inst_nospace.isin({v for v in inst_values if inst in v or v in inst})
        hit = std[mask].copy()
        if hit.empty:
            log.info("시설 %s '%s': 재검색 결과 없음", fac.get("시설ID"), fac["검수_시설명"])
            continue
        if len(hit) > max_hits_warn:
            log.warning("시설 %s '%s': %d건 매칭 — 검수_시설명이 너무 일반적인지 확인", fac.get("시설ID"), fac["검수_시설명"], len(hit))
        hit["시설ID"] = fac["시설ID"]
        hit["시설명"] = fac["검수_시설명"]
        hit["사업유형"] = [work_type_for(n, str(fac.get("검색어") or ""), fac["검수_시설명"]) for n in hit["공고명"]]
        hit["표2_대분류"] = fac["표2_대분류"]
        hit["표2_중분류"] = fac["표2_중분류"]
        hit["검색어"] = fac["검색어"]
        hit["수요기관_대표"] = fac.get("검수_수요기관") or fac.get("수요기관", "")
        out.append(hit)
    df = pd.concat(out, ignore_index=True) if out else pd.DataFrame()
    if not df.empty:
        # 유지보수 성격 공고는 시설 재검색에서도 제외 (신축·증축·리모델링 프로젝트 공사비만)
        df = df[~df["사업유형"].isin(["유지보수"])].copy()
        df = df.drop_duplicates(["시설ID", "공고키"])
    log.info("시설 재검색: %d개 시설 → %d건 공고", reviewed["시설ID"].nunique() if not reviewed.empty else 0, len(df))
    return df
