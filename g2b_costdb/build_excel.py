"""S7 Excel DB 생성 (openpyxl, 수식 기반).

시트: README / 01_시설마스터 / 02_공고목록_전체 / 03_공고목록_최신 / 04_공사비DB_공종별 / 05_공사비DB_시설합산 / 06_검증로그 / 07_추출노트
색상: 파란 글자 = 원천값(API/문서), 검은 글자 = 수식, 노란 배경 = 사용자 입력·검수 셀
"""
from __future__ import annotations

import datetime as dt
import json
import math
from typing import Dict, List, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = Font(name="Arial", size=10)
BLUE = Font(name="Arial", size=10, color="0000FF")
BOLD = Font(name="Arial", size=10, bold=True)
HEAD = PatternFill("solid", fgColor="DDEBF7")
YELLOW = PatternFill("solid", fgColor="FFFF00")
MONEY = "#,##0;(#,##0);-"
TRADES = ["건축", "전기", "정보통신", "소방", "조경", "기계설비", "토목", "기타"]

S_TRADE = "04_공사비DB_공종별"
S_FAC = "05_공사비DB_시설합산"


def _clean(v):
    if v is None or v is pd.NA or v is pd.NaT:
        return None
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return None
    if isinstance(v, (pd.Timestamp, dt.datetime)):
        return v.to_pydatetime() if isinstance(v, pd.Timestamp) else v
    if hasattr(v, "item"):
        try:
            return v.item()
        except (ValueError, TypeError):
            return str(v)
    if isinstance(v, str) and v in ("nan", "None", "<NA>", "NaT"):
        return ""
    if isinstance(v, (dict, list, tuple, set)):
        return json.dumps(v, ensure_ascii=False, default=str)
    return v


def _write_df(ws, df: pd.DataFrame, money_cols: Optional[List[str]] = None, blue_cols: Optional[List[str]] = None,
              yellow_cols: Optional[List[str]] = None, widths: Optional[Dict[str, int]] = None) -> None:
    cols = list(df.columns)
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font, cell.fill = BOLD, HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in df.itertuples(index=False):
        ws.append([_clean(v) for v in row])
    money_idx = [cols.index(c) + 1 for c in (money_cols or []) if c in cols]
    blue_idx = [cols.index(c) + 1 for c in (blue_cols or []) if c in cols]
    yellow_idx = [cols.index(c) + 1 for c in (yellow_cols or []) if c in cols]
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BLUE if c in blue_idx else FONT
            if c in money_idx:
                cell.number_format = MONEY
            if c in yellow_idx:
                cell.fill = YELLOW
    for i, c in enumerate(cols, 1):
        w = (widths or {}).get(c, min(40, max(10, len(str(c)) * 1.6 + 4)))
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


NOTICE_COLS = ["공고번호", "공고차수", "공고키", "공고명", "공고종류", "재공고여부", "공고일시", "입찰마감일시", "개찰일시", "공고기관",
               "수요기관", "공사현장지역", "추정가격", "기초금액", "주공종명", "면허제한업종", "공종", "공종근거", "사업유형", "사업유형_원분류",
               "시설ID", "시설명", "프로젝트ID", "프로젝트키", "최신여부", "대체공고번호", "대표선정사유", "사전규격번호", "첨부파일수", "상세URL"]


def build_workbook(path: str, facility: pd.DataFrame, hist: pd.DataFrame, latest: pd.DataFrame,
                   trade: pd.DataFrame, logs: pd.DataFrame, notes: pd.DataFrame, meta: Optional[Dict] = None) -> None:
    wb = Workbook()
    # ── README ──
    ws = wb.active
    ws.title = "README"
    meta = meta or {}
    lines = [
        ["LIMAC 공사비 DB — 나라장터 발주공고 기반"],
        ["생성일", dt.date.today().isoformat()],
        ["수집기간", meta.get("period", "")],
        ["출처", "조달청 나라장터 입찰공고정보서비스(공공데이터포털) + 공고 첨부문서(공고문·현장설명서 등) LLM 추출"],
        ["금액 기준", "낙찰가 아님. 추정가격(VAT 제외)·기초금액(추정가격+VAT) = 예정가격 산정 기준. 관급자재는 별도 컬럼."],
        ["대표 공고", "동일 (프로젝트, 공종)에 복수 공고 시 취소공고 제외 후 최신 공고 1건을 대표로 채택(변경·재공고 반영). 이력은 02시트 보존."],
        ["프로젝트ID", "시설ID-사업유형코드(N 신축 / E 증축 / R 리모델링 / ER 증축·리모델링). 같은 시설의 신축과 리모델링은 별도 프로젝트로 집계되며, ㎡당 공사비 비교는 반드시 사업유형이 같은 프로젝트끼리 할 것."],
        ["색상", "파란 글자=원천값(API/문서), 검은 글자=수식, 노란 배경=사용자 입력·검수 셀 (01 시트의 연면적·층수 등은 문서 추출값이며 직접 고치면 05 ㎡당 공사비에 반영됨)"],
        ["단위", "금액: 원 / 면적: ㎡ / 기간: 일"],
        [],
        ["시트", "내용"],
        ["01_시설마스터", "시설(프로젝트) 단위 개요·공종 커버리지"],
        ["02_공고목록_전체", "수집된 모든 공고(차수·재공고·취소 이력 포함)"],
        ["03_공고목록_최신", "프로젝트키(시설×공종)당 대표 공고"],
        ["04_공사비DB_공종별", "대표 공고 기준 공종별 금액 구성(API+문서), 총공사비 수식 (관급자재는 API 값 우선, 없으면 문서 추출값)"],
        ["05_공사비DB_시설합산", "시설 단위 공종 합산(SUMIFS), ㎡당 공사비, 보정계수 입력, 공종 누락 경고"],
        ["06_검증로그", "API-문서 교차검증, 재발주 금액변동, 정합성 점검 결과"],
        ["07_추출노트", "첨부파일별 다운로드·텍스트 추출 결과"],
    ]
    for l in lines:
        ws.append(l)
    ws["A1"].font = Font(name="Arial", size=12, bold=True)
    for r in range(2, ws.max_row + 1):
        for c in (1, 2):
            ws.cell(row=r, column=c).font = FONT
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 22, 110

    # ── 01 시설마스터 ──
    ws = wb.create_sheet("01_시설마스터")
    fac = facility.copy()
    for c in ["연면적_m2", "건축면적_m2", "지하층수", "지상층수", "구조", "용도", "공사기간_일", "검수_포함여부", "메모"]:
        if c not in fac.columns:
            fac[c] = "포함" if c == "검수_포함여부" else None
    _write_df(ws, fac, money_cols=["연면적_m2", "건축면적_m2"],
              blue_cols=[c for c in fac.columns if c not in ("검수_포함여부", "메모")],
              yellow_cols=["연면적_m2", "건축면적_m2", "지하층수", "지상층수", "구조", "용도", "공사기간_일", "검수_포함여부", "메모"],
              widths={"시설명": 34, "공고명_예시": 50, "수요기관": 24})
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=list(fac.columns).index("연면적_m2") + 1).number_format = "#,##0.0"

    # ── 02 / 03 공고목록 ──
    for title, df in (("02_공고목록_전체", hist), ("03_공고목록_최신", latest)):
        ws = wb.create_sheet(title)
        cols = [c for c in NOTICE_COLS if c in df.columns]
        _write_df(ws, df[cols], money_cols=["추정가격", "기초금액"], blue_cols=cols,
                  widths={"공고명": 60, "수요기관": 24, "공고기관": 24, "상세URL": 40, "프로젝트키": 26})

    # ── 04 공종별 ──
    ws = wb.create_sheet(S_TRADE)
    tcols = ["프로젝트ID", "시설ID", "시설명", "사업유형", "공종", "공고번호", "공고차수", "공고명", "공고일시", "수요기관",
             "추정가격_API", "기초금액_API", "부가세(수식)", "관급자재_API", "도급자관급액_API", "관급자관급액_API",
             "도급자관급액_문서", "관급자관급액_문서", "총공사비(수식)",
             "예산금액_API", "공사기간_일_문서", "추정가격_문서", "기초금액_문서", "추정가격_차이율(수식)", "신뢰도", "근거문구", "출처파일", "상세URL"]
    t = trade.copy()
    for c in tcols:
        if c not in t.columns:
            t[c] = None
    t = t[tcols]
    _write_df(ws, t, money_cols=["추정가격_API", "기초금액_API", "부가세(수식)", "관급자재_API", "도급자관급액_API", "관급자관급액_API",
                                 "도급자관급액_문서", "관급자관급액_문서", "총공사비(수식)", "예산금액_API", "추정가격_문서", "기초금액_문서"],
              blue_cols=[c for c in tcols if "(수식)" not in c],
              widths={"시설명": 30, "공고명": 55, "근거문구": 50, "출처파일": 30, "상세URL": 40})
    col = {c: get_column_letter(i + 1) for i, c in enumerate(tcols)}
    for r in range(2, ws.max_row + 1):
        P, B = f"{col['추정가격_API']}{r}", f"{col['기초금액_API']}{r}"
        Ga, G1a, G2a = f"{col['관급자재_API']}{r}", f"{col['도급자관급액_API']}{r}", f"{col['관급자관급액_API']}{r}"
        G1, G2 = f"{col['도급자관급액_문서']}{r}", f"{col['관급자관급액_문서']}{r}"
        Pd, Bd = f"{col['추정가격_문서']}{r}", f"{col['기초금액_문서']}{r}"
        ws[f"{col['부가세(수식)']}{r}"] = f'=IF(AND(ISNUMBER({B}),ISNUMBER({P})),{B}-{P},IF(AND(ISNUMBER({Bd}),ISNUMBER({Pd})),{Bd}-{Pd},""))'
        # 총공사비 = [기초금액(API) → 추정가격(API)×1.1 → 기초금액(문서) → 추정가격(문서)×1.1 → 0]
        #          + 관급자재 [API 도급자+관급자 → API 합계 → 문서 도급자+관급자]
        base = f'IF(ISNUMBER({B}),{B},IF(ISNUMBER({P}),{P}*1.1,IF(ISNUMBER({Bd}),{Bd},IF(ISNUMBER({Pd}),{Pd}*1.1,0))))'
        gov = f'IF(N({G1a})+N({G2a})>0,N({G1a})+N({G2a}),IF(N({Ga})>0,N({Ga}),N({G1})+N({G2})))'
        ws[f"{col['총공사비(수식)']}{r}"] = f"={base}+{gov}"
        ws[f"{col['추정가격_차이율(수식)']}{r}"] = f'=IF(AND(ISNUMBER({P}),ISNUMBER({Pd}),{P}<>0),({Pd}-{P})/{P},"")'
        ws[f"{col['추정가격_차이율(수식)']}{r}"].number_format = "0.00%"
        for c in ("부가세(수식)", "총공사비(수식)", "추정가격_차이율(수식)"):
            ws[f"{col[c]}{r}"].font = FONT

    # ── 05 시설합산 ──
    ws = wb.create_sheet(S_FAC)
    fcols = ["프로젝트ID", "시설ID", "시설명", "사업유형", "표2_대분류", "표2_중분류", "수요기관", "연면적_m2", "지상층수", "지하층수", "구조",
             "공고연도"] + [f"{tr}(수식)" for tr in TRADES] + ["총공사비합계(수식)", "㎡당공사비_원(수식)", "건설공사비지수_보정계수(입력)",
                                                              "보정_㎡당공사비_원(수식)", "공종누락경고(수식)", "공고건수_최신", "메모"]
    f5 = pd.DataFrame({c: facility[c] if c in facility.columns else None for c in
                       ["프로젝트ID", "시설ID", "시설명", "사업유형", "표2_대분류", "표2_중분류", "수요기관", "연면적_m2", "지상층수", "지하층수", "구조",
                        "공고건수_최신"]}, index=facility.index)
    f5["공고연도"] = facility["최종공고일"].astype(str).str[:4] if ("최종공고일" in facility.columns and len(facility)) else None
    f5["건설공사비지수_보정계수(입력)"] = 1.0
    f5["메모"] = ""
    for c in fcols:
        if c not in f5.columns:
            f5[c] = None
    f5 = f5[fcols]
    _write_df(ws, f5, money_cols=["연면적_m2"] + [f"{tr}(수식)" for tr in TRADES] + ["총공사비합계(수식)", "㎡당공사비_원(수식)",
                                                                              "보정_㎡당공사비_원(수식)"],
              blue_cols=["프로젝트ID", "시설ID", "시설명", "사업유형", "표2_대분류", "표2_중분류", "수요기관", "연면적_m2", "지상층수", "지하층수", "구조",
                         "공고연도", "공고건수_최신"],
              yellow_cols=["건설공사비지수_보정계수(입력)", "메모"], widths={"시설명": 30, "공종누락경고(수식)": 28})
    fc = {c: get_column_letter(i + 1) for i, c in enumerate(fcols)}
    tot_rng = f"'{S_TRADE}'!${col['총공사비(수식)']}:${col['총공사비(수식)']}"
    id_rng, tr_rng = f"'{S_TRADE}'!${col['프로젝트ID']}:${col['프로젝트ID']}", f"'{S_TRADE}'!${col['공종']}:${col['공종']}"
    # 05 연면적은 01_시설마스터(노란 셀, 사용자 수정 가능)를 참조 → 01 에서 고치면 ㎡당 공사비에 반영
    fac_cols = list(fac.columns)
    area_col = get_column_letter(fac_cols.index("연면적_m2") + 1)
    n_fac = max(len(fac), 1) + 1                     # 열 전체($V:$V) 대신 행 범위를 지정(수식 검증 도구·구형 Excel 호환)
    fac_rng = f"'01_시설마스터'!${area_col}$2:${area_col}${n_fac}"
    fac_id_rng = f"'01_시설마스터'!$A$2:$A${n_fac}"
    for r in range(2, ws.max_row + 1):
        for tr in TRADES:
            ws[f"{fc[f'{tr}(수식)']}{r}"] = f'=SUMIFS({tot_rng},{id_rng},$A{r},{tr_rng},"{tr}")'
        first, last = fc[f"{TRADES[0]}(수식)"], fc[f"{TRADES[-1]}(수식)"]
        ws[f"{fc['총공사비합계(수식)']}{r}"] = f"=SUM({first}{r}:{last}{r})"
        A, T = f"{fc['연면적_m2']}{r}", f"{fc['총공사비합계(수식)']}{r}"
        ws[A] = f'=IFERROR(INDEX({fac_rng},MATCH($A{r},{fac_id_rng},0)),"")'
        ws[A].font = FONT
        ws[f"{fc['㎡당공사비_원(수식)']}{r}"] = f'=IF(AND(ISNUMBER({A}),{A}>0),{T}/{A},"")'
        ws[f"{fc['보정_㎡당공사비_원(수식)']}{r}"] = (f'=IF(ISNUMBER({fc["㎡당공사비_원(수식)"]}{r}),'
                                                 f'{fc["㎡당공사비_원(수식)"]}{r}*{fc["건설공사비지수_보정계수(입력)"]}{r},"")')
        # 공고 자체가 없으면 '누락', 공고는 있는데 금액이 0 이면 '금액없음'
        miss = "&".join(f'IF(COUNTIFS({id_rng},$A{r},{tr_rng},"{tr}")=0,"{tr} ","")' for tr in ("건축", "전기", "정보통신", "소방"))
        zero = "&".join(f'IF(AND(COUNTIFS({id_rng},$A{r},{tr_rng},"{tr}")>0,{fc[f"{tr}(수식)"]}{r}=0),"{tr} ","")'
                        for tr in ("건축", "전기", "정보통신", "소방"))
        ws[f"{fc['공종누락경고(수식)']}{r}"] = (f'=TRIM(IF(TRIM({miss})="","",TRIM({miss})&"누락 ")'
                                            f'&IF(TRIM({zero})="","",TRIM({zero})&"금액없음"))')
        for c in fcols:
            if "(수식)" in c:
                ws[f"{fc[c]}{r}"].font = FONT
    n = ws.max_row + 2
    ws[f"A{n}"] = ("주: 건설공사비지수 보정계수는 사용자 입력(기준연도 지수/공고연도 지수). 총공사비 = 기초금액(API, 없으면 추정가격×1.1, 그것도 없으면 문서 추출값) + 관급자재(API 도급자·관급자 설치액 우선, 없으면 API 합계, 없으면 문서 추출값). "
                   "연면적은 01_시설마스터 값을 참조(01 에서 수정). 공종누락경고: '누락'=해당 공종 공고 없음, '금액없음'=공고는 있으나 금액 0. "
                   "㎡당 공사비는 사업유형(신축/증축/리모델링)이 같은 프로젝트끼리만 비교할 것 — 리모델링의 연면적은 공사 대상 연면적임.")
    ws[f"A{n}"].font = FONT

    # ── 06 / 07 ──
    ws = wb.create_sheet("06_검증로그")
    lcols = ["공고번호", "항목", "API값", "문서값", "차이", "차이율", "판정", "비고"]
    lg = logs.copy() if not logs.empty else pd.DataFrame(columns=lcols)
    for c in lcols:
        if c not in lg.columns:
            lg[c] = None
    _write_df(ws, lg[lcols], money_cols=["API값", "문서값", "차이"], blue_cols=lcols, widths={"비고": 70, "항목": 26})
    for r in range(2, ws.max_row + 1):
        ws[f"F{r}"].number_format = "0.00%"
    ws = wb.create_sheet("07_추출노트")
    ncols = ["공고번호", "파일명", "형식", "우선순위", "다운로드", "추출성공", "추출글자수", "파서", "LLM호출", "오류", "URL"]
    nt = notes.copy() if not notes.empty else pd.DataFrame(columns=ncols)
    for c in ncols:
        if c not in nt.columns:
            nt[c] = None
    _write_df(ws, nt[ncols], blue_cols=ncols, widths={"파일명": 45, "오류": 50, "URL": 50})

    wb.save(path)
