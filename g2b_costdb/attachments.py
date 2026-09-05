"""S5 첨부파일 다운로드 + 텍스트 추출.

지원: HWP(5.0 바이너리, olefile+zlib 레코드 파서 → hwp5txt → 한컴 COM 폴백) / HWPX / PDF / DOCX / XLSX / ZIP
결과: data/files/<공고번호>/<파일명>, data/text/<공고번호>__<파일명>.txt, 추출노트 rows
"""
from __future__ import annotations

import io
import logging
import os
import re
import shutil
import struct
import subprocess
import zipfile
import zlib
from typing import Dict, List, Optional, Tuple
from urllib.parse import unquote

import requests

log = logging.getLogger(__name__)

# 파일 선별 우선순위 (낮을수록 먼저). 매칭 없으면 90, 제외 확장자는 None
_PRIORITY = [
    (r"공고문|입찰공고|공고서", 10),
    (r"현장설명|공사설명|설계설명|사업설명|공사개요", 20),
    (r"관급|자재", 30),
    (r"내역서|산출내역|원가계산|총괄", 40),
    (r"특수조건|시방|과업", 60),
]
_SKIP_EXT = {".dwg", ".dxf", ".jpg", ".jpeg", ".png", ".gif", ".bmp", ".exe", ".msi", ".tif", ".tiff", ".mp4"}


def file_priority(filename: str) -> Optional[int]:
    ext = os.path.splitext(filename.lower())[1]
    if ext in _SKIP_EXT:
        return None
    for pat, pr in _PRIORITY:
        if re.search(pat, filename):
            return pr
    return 90


_EMPTY = {"", "nan", "none", "null"}


def safe_name(s: str) -> str:
    """Windows 금지문자 제거 + 길이 제한(확장자는 보존)."""
    s = unquote(str(s or "")).strip().replace("/", "_").replace("\\", "_")
    s = re.sub(r'[<>:"|?*\x00-\x1f]', "_", s)
    if s.lower() in _EMPTY:
        return ""
    stem, ext = os.path.splitext(s)
    return (stem[:100] + ext[:10]) if stem else s[:110]


_MAGIC = [(b"%PDF", ".pdf"), (b"\xd0\xcf\x11\xe0", ".hwp"), (b"PK\x03\x04", ".zip")]


def sniff_ext(head: bytes) -> str:
    """앞부분 바이트로 확장자 추정. ZIP 계열은 내부 파일명으로 HWPX/DOCX/XLSX 를 구분."""
    for magic, ext in _MAGIC:
        if head.startswith(magic):
            if ext == ".zip":
                if b"Contents/section" in head or b"application/hwp+zip" in head:
                    return ".hwpx"
                if b"word/document.xml" in head or b"word/_rels" in head:
                    return ".docx"
                if b"xl/workbook.xml" in head or b"xl/_rels" in head:
                    return ".xlsx"
            return ext
    return ""


def looks_like_html(head: bytes, content_type: str = "") -> bool:
    h = head.lstrip()[:64].lower()
    return "text/html" in (content_type or "").lower() or h.startswith(b"<!doctype") or h.startswith(b"<html") or h.startswith(b"<script")


# ── 다운로드 ─────────────────────────────────────────────────────
def _cd_filename(cd: str) -> str:
    m = re.search(r"filename\*=(?:UTF-8'')?\"?([^\";]+)", cd) or re.search(r"filename=\"?([^\";]+)", cd)
    if not m:
        return ""
    name = m.group(1).strip()
    try:
        name = name.encode("latin-1").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        try:
            name = name.encode("latin-1").decode("cp949")
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    return unquote(name)


def download(url: str, dest_dir: str, hint_name: str = "", timeout: int = 120) -> Optional[str]:
    """첨부 1개 다운로드 → 저장 경로(실패 시 None). 이미 받은 파일은 요청 없이 재사용, 임시파일(.part)로 받아 완료 후 교체,
    HTML(로그인·오류 페이지) 응답은 실패로 처리, 확장자가 없으면 Content-Disposition/매직바이트로 보완."""
    os.makedirs(dest_dir, exist_ok=True)
    hint = safe_name(hint_name)
    if hint and os.path.splitext(hint)[1]:
        path = os.path.join(dest_dir, hint)
        if os.path.exists(path) and os.path.getsize(path) > 0:
            return path
    try:
        with requests.get(url, stream=True, timeout=timeout,
                          headers={"User-Agent": "Mozilla/5.0 (LIMAC cost-db collector)", "Referer": "https://www.g2b.go.kr/"}) as r:
            r.raise_for_status()
            it = r.iter_content(1 << 16)
            first = b""
            for chunk in it:
                if chunk:
                    first = chunk
                    break
            if looks_like_html(first, r.headers.get("Content-Type", "")):
                log.warning("다운로드 실패(HTML 응답 — 로그인/오류 페이지) %s", url)
                return None
            name = hint or safe_name(_cd_filename(r.headers.get("Content-Disposition", ""))) or safe_name(url.split("/")[-1].split("?")[0])
            if not os.path.splitext(name)[1]:
                ext = sniff_ext(first) or os.path.splitext(safe_name(_cd_filename(r.headers.get("Content-Disposition", ""))))[1]
                name = (name or "file") + ext
            path = os.path.join(dest_dir, name)
            if os.path.exists(path) and os.path.getsize(path) > 0:
                return path
            tmp = path + ".part"
            with open(tmp, "wb") as f:
                if first:
                    f.write(first)
                for chunk in it:
                    if chunk:
                        f.write(chunk)
            if os.path.getsize(tmp) == 0:
                os.remove(tmp)
                log.warning("다운로드 실패(빈 파일) %s", url)
                return None
            os.replace(tmp, path)
            return path
    except (requests.RequestException, OSError) as e:
        log.warning("다운로드 실패 %s: %s", url, e)
        try:
            if "tmp" in locals() and os.path.exists(tmp):
                os.remove(tmp)
        except OSError:
            pass
        return None


# ── HWP 5.0 (olefile) ────────────────────────────────────────────
_HWPTAG_PARA_TEXT = 0x10 + 51
_CTRL_8 = {1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23}


def _hwp_records(data: bytes):
    pos, n = 0, len(data)
    while pos + 4 <= n:
        hdr = struct.unpack_from("<I", data, pos)[0]
        tag, size = hdr & 0x3FF, (hdr >> 20) & 0xFFF
        pos += 4
        if size == 0xFFF:
            size = struct.unpack_from("<I", data, pos)[0]
            pos += 4
        yield tag, data[pos: pos + size]
        pos += size


def _hwp_para_text(payload: bytes) -> str:
    out, i, n = [], 0, len(payload) - (len(payload) % 2)
    while i + 1 < n:
        c = payload[i] | (payload[i + 1] << 8)
        if c < 32:
            if c in _CTRL_8:
                out.append("\t" if c == 9 else "")
                i += 16
                continue
            if c in (10, 13):
                out.append("\n")
            elif c in (30, 31):
                out.append(" ")
            i += 2
            continue
        out.append(chr(c))
        i += 2
    return "".join(out)


def extract_hwp(path: str) -> Tuple[str, str]:
    """반환 (텍스트, 파서명). 배포용/암호화 문서는 예외."""
    import olefile  # type: ignore

    ole = olefile.OleFileIO(path)
    try:
        header = ole.openstream("FileHeader").read()
        flags = struct.unpack_from("<I", header, 36)[0]
        compressed, encrypted, distributed = flags & 1, (flags >> 1) & 1, (flags >> 2) & 1
        if encrypted:
            raise ValueError("암호화된 HWP")
        if distributed:
            raise ValueError("배포용(DRM) HWP — hwp5txt/한컴 COM 폴백 필요")
        sections = sorted([e for e in ole.listdir() if e[0] == "BodyText"],
                          key=lambda e: int(re.sub(r"\D", "", e[1]) or 0))
        texts = []
        for e in sections:
            raw = ole.openstream(e).read()
            data = zlib.decompress(raw, -15) if compressed else raw
            for tag, payload in _hwp_records(data):
                if tag == _HWPTAG_PARA_TEXT:
                    t = _hwp_para_text(payload)
                    if t.strip():
                        texts.append(t)
        text = "\n".join(texts)
        if len(text.strip()) < 50:
            raise ValueError("본문 텍스트 미검출")
        return text, "olefile"
    finally:
        ole.close()


def extract_hwp_with_fallbacks(path: str) -> Tuple[str, str]:
    try:
        return extract_hwp(path)
    except Exception as e:  # noqa: BLE001
        log.info("olefile 파서 실패(%s) → hwp5txt 시도", e)
    if shutil.which("hwp5txt"):
        try:
            out = subprocess.run(["hwp5txt", path], capture_output=True, timeout=180)
            txt = out.stdout.decode("utf-8", errors="ignore")
            if len(txt.strip()) > 50:
                return txt, "hwp5txt"
        except (subprocess.SubprocessError, OSError) as e:
            log.info("hwp5txt 실패: %s", e)
    if os.name == "nt":
        try:
            return _extract_hwp_com(path), "hancom_com"
        except Exception as e:  # noqa: BLE001
            log.info("한컴 COM 실패: %s", e)
    raise ValueError("HWP 텍스트 추출 실패(모든 백엔드)")


def _extract_hwp_com(path: str) -> str:
    """Windows + 한컴오피스 설치 환경에서 HWP → TXT (pywin32 필요)."""
    import win32com.client  # type: ignore

    hwp = win32com.client.Dispatch("HWPFrame.HwpObject")
    hwp.RegisterModule("FilePathCheckDLL", "FilePathCheckerModule")
    hwp.Open(os.path.abspath(path), "HWP", "forceopen:true")
    tmp = os.path.abspath(path) + ".txt"
    hwp.SaveAs(tmp, "TEXT")
    hwp.Quit()
    with open(tmp, encoding="utf-8", errors="ignore") as f:
        return f.read()


# ── HWPX / PDF / DOCX / XLSX ─────────────────────────────────────
def extract_hwpx(path: str) -> str:
    from xml.etree import ElementTree as ET
    ns_p = "{http://www.hancom.co.kr/hwpml/2011/paragraph}"
    lines: List[str] = []
    with zipfile.ZipFile(path) as z:
        secs = sorted([n for n in z.namelist() if re.match(r"Contents/section\d+\.xml", n)],
                      key=lambda n: int(re.sub(r"\D", "", n)))
        for n in secs:
            root = ET.fromstring(z.read(n))

            def ptext(p):
                out = []
                for t in p.iter(ns_p + "t"):
                    out.append(t.text or "")
                    for c in t:                       # <hp:tab/>, <hp:lineBreak/> 뒤의 tail 텍스트 보존
                        tag = c.tag.split("}")[-1]
                        out.append("\t" if tag == "tab" else ("\n" if tag == "lineBreak" else ""))
                        out.append(c.tail or "")
                return "".join(out)

            def walk(el):
                tag = el.tag.split("}")[-1]
                if tag == "tbl":
                    for tr in el.findall(ns_p + "tr"):
                        cells = ["/".join(ptext(p) for p in tc.findall(".//" + ns_p + "p")).strip()
                                 for tc in tr.findall(ns_p + "tc")]
                        lines.append("\t".join(cells))
                    return
                if tag == "p" and el.find(".//" + ns_p + "tbl") is None:
                    t = ptext(el)
                    if t.strip():
                        lines.append(t)
                    return
                for c in el:
                    walk(c)

            walk(root)
    return "\n".join(lines)


def extract_pdf(path: str) -> Tuple[str, str]:
    import pdfplumber  # type: ignore

    parts: List[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            parts.append(page.extract_text() or "")
            for tb in page.extract_tables() or []:
                for row in tb:
                    parts.append("\t".join((c or "").replace("\n", " ") for c in row))
    text = "\n".join(parts)
    if len(text.strip()) > 100:
        return text, "pdfplumber"
    # 스캔 PDF → OCR (선택 의존성)
    try:
        import pytesseract  # type: ignore
        from pdf2image import convert_from_path  # type: ignore
        ocr = [pytesseract.image_to_string(img, lang="kor+eng") for img in convert_from_path(path, dpi=200)]
        return "\n".join(ocr), "ocr"
    except Exception as e:  # noqa: BLE001
        raise ValueError(f"PDF 텍스트 없음(스캔) & OCR 불가: {e}")


def extract_docx(path: str) -> str:
    import docx  # type: ignore
    d = docx.Document(path)
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows:
            lines.append("\t".join(c.text.strip() for c in r.cells))
    return "\n".join(lines)


def extract_xlsx(path: str, max_rows: int = 400) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"## {ws.title}")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            if any(v is not None for v in row):
                lines.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(lines)


def extract_any(path: str) -> Tuple[str, str]:
    ext = os.path.splitext(path.lower())[1]
    if ext == ".hwp":
        return extract_hwp_with_fallbacks(path)
    if ext == ".hwpx":
        return extract_hwpx(path), "hwpx-xml"
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".docx":
        return extract_docx(path), "python-docx"
    if ext in (".xlsx", ".xlsm"):
        return extract_xlsx(path), "openpyxl"
    if ext == ".zip":
        return extract_zip(path)
    if ext in (".txt", ".csv"):
        raw = open(path, "rb").read()
        try:
            return raw.decode("utf-8"), "text"
        except UnicodeDecodeError:
            return raw.decode("cp949", errors="replace"), "text(cp949)"
    raise ValueError(f"미지원 형식: {ext}")


def extract_zip(path: str) -> Tuple[str, str]:
    out, parsers = [], []
    tmp = path + "_unz"
    os.makedirs(tmp, exist_ok=True)
    with zipfile.ZipFile(path) as z:
        members = []
        for info in z.infolist():
            if info.is_dir():
                continue
            name = info.filename
            if not (info.flag_bits & 0x800):        # UTF-8 플래그가 없으면 CP949 로 저장된 한글 파일명 복원
                try:
                    name = info.filename.encode("cp437").decode("cp949")
                except (UnicodeEncodeError, UnicodeDecodeError):
                    pass
            pr = file_priority(name)
            if pr is not None:
                members.append((pr, info, name))
        for pr, info, name in sorted(members, key=lambda x: x[0])[:6]:
            dest = os.path.join(tmp, safe_name(os.path.basename(name)))
            with z.open(info) as src, open(dest, "wb") as dst:
                shutil.copyfileobj(src, dst)
            try:
                t, p = extract_any(dest)
                out.append(f"### [{name}]\n{t}")
                parsers.append(p)
            except Exception as e:  # noqa: BLE001
                log.info("zip 내부 %s 추출 실패: %s", name, e)
    if not out:
        raise ValueError("zip 내 추출 가능 파일 없음")
    return "\n\n".join(out), "zip(" + ",".join(sorted(set(parsers))) + ")"


# ── 공고 단위 처리 ───────────────────────────────────────────────
def process_notice_attachments(row: Dict, files_dir: str, text_dir: str, attach_max: int = 10,
                               max_files: int = 4) -> Tuple[str, List[Dict]]:
    """한 공고의 첨부 중 우선순위 상위 파일을 내려받아 텍스트 추출. 반환 (결합 텍스트, 추출노트 rows)."""
    bid_no = str(row["공고번호"])
    cands = []
    for i in range(1, attach_max + 1):
        url, name = str(row.get(f"첨부URL{i}") or "").strip(), str(row.get(f"첨부파일명{i}") or "").strip()
        if len(url) < 5 or url.lower() in _EMPTY:
            continue
        if name.lower() in _EMPTY:
            name = ""
        pr = file_priority(name or url)
        if pr is not None:
            cands.append((pr, url, name))
    cands.sort(key=lambda x: x[0])
    notes, texts = [], []
    os.makedirs(text_dir, exist_ok=True)
    for pr, url, name in cands[:max_files]:
        note = {"공고번호": bid_no, "파일명": name, "우선순위": pr, "URL": url, "다운로드": "N", "추출성공": "N",
                "추출글자수": 0, "파서": "", "오류": ""}
        path = download(url, os.path.join(files_dir, bid_no), hint_name=name)
        if not path:
            note["오류"] = "다운로드 실패"
            notes.append(note)
            continue
        note["다운로드"] = "Y"
        note["형식"] = os.path.splitext(path)[1].lower()
        try:
            text, parser = extract_any(path)
            note.update({"추출성공": "Y", "추출글자수": len(text), "파서": parser})
            tpath = os.path.join(text_dir, f"{bid_no}__{safe_name(os.path.basename(path))}.txt")
            with open(tpath, "w", encoding="utf-8") as f:
                f.write(text)
            texts.append(f"===== [{os.path.basename(path)}] =====\n{text}")
        except Exception as e:  # noqa: BLE001
            note["오류"] = str(e)[:200]
        notes.append(note)
    return "\n\n".join(texts), notes
