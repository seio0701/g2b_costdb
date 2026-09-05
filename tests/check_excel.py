"""Excel 산출물의 수식 오류 점검. 실행: python -m tests.check_excel output/공사비DB.xlsx
LibreOffice(soffice)가 있으면 재계산본을 만들어 읽고, 없으면 `pip install formulas` 로 파이썬에서 수식을 평가한다."""
from __future__ import annotations

import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl


def _errors_in(path: str) -> list:
    wb = openpyxl.load_workbook(path, data_only=True)
    errs = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#"):
                    errs.append((ws.title, c.coordinate, c.value))
    return errs


def check(path: str) -> int:
    n_formula = sum(1 for ws in openpyxl.load_workbook(path).worksheets for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("="))
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run([soffice, "--headless", "--convert-to", "xlsx", "--outdir", tmp, os.path.abspath(path)],
                           capture_output=True, timeout=300)
            out = os.path.join(tmp, os.path.basename(path))
            if os.path.exists(out):
                errs = _errors_in(out)
                print(f"LibreOffice 재계산: 수식 {n_formula}개, 오류 셀 {len(errs)}개 {errs[:10]}")
                return 1 if errs else 0
            print("LibreOffice 재계산 실패(Calc 미설치 등) → formulas 패키지로 평가")
    try:
        import formulas  # type: ignore
    except ImportError:
        print("formulas 패키지가 없습니다: pip install formulas  (또는 Excel 에서 파일을 열어 오류 셀(#) 유무를 확인)")
        return 2
    import logging, warnings
    warnings.filterwarnings("ignore"); logging.disable(logging.CRITICAL)
    sol = formulas.ExcelModel().loads(path).finish().calculate()
    bad = 0
    for k, v in sol.items():
        val = getattr(v, "value", v)
        sv = val.tolist() if hasattr(val, "tolist") else val
        while isinstance(sv, list) and sv:
            sv = sv[0]
        if isinstance(sv, str) and sv.startswith("#"):
            bad += 1
            if bad <= 10:
                print("오류 셀:", k, sv)
    print(f"formulas 평가: 수식 {n_formula}개, 오류 {bad}개")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(check(sys.argv[1] if len(sys.argv) > 1 else os.path.join("output", "공사비DB.xlsx")))
