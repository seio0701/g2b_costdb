"""가짜 나라장터 서버로 전 단계 통합 실행: probe → collect → discover → research → dedup → attach → extract(견적/--yes 모의) → excel.
실행: python -m tests.test_e2e_mock   (네트워크·API 키 불필요, 임시 폴더에서 동작)
※ 모든 데이터는 가상의 예시. LLM 호출은 텍스트에서 정규식으로 값을 뽑는 가짜 함수로 대체한다.
"""
from __future__ import annotations

import contextlib
import io
import json
import os
import re
import sys
import tempfile
from unittest import mock

import pandas as pd
import yaml

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from g2b_costdb import extract_llm, pipeline  # noqa: E402
from g2b_costdb.classify import ROOT, load_config  # noqa: E402
from tests.mock_g2b import KEY, MockG2B  # noqa: E402


def _run(argv):
    """pipeline.main 을 실행하고 표준출력을 돌려준다."""
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        pipeline.main(argv)
    out = buf.getvalue()
    print(out[-1500:] if len(out) > 1500 else out)
    return out


def _fake_extract(text: str, api_hint: dict, cfg_llm: dict, max_tokens=None) -> dict:
    """LLM 대신 정규식으로 공고문 텍스트에서 값을 뽑는 모의 추출기."""
    def amt(label):
        m = re.search(label + r"[^\d]{0,20}([\d,]{5,})", text)
        return int(m.group(1).replace(",", "")) if m else None
    area = re.search(r"연면적\s*([\d,\.]+)\s*㎡", text)
    floors = re.search(r"지상\s*(\d+)\s*층", text)
    days = re.search(r"(\d[\d,]*)\s*일", text)
    doc = {"공사명": api_hint.get("공고명"), "사업유형": "신축", "공종": "건축",
           "연면적_m2": float(area.group(1).replace(",", "")) if area else None,
           "지상층수": int(floors.group(1)) if floors else None,
           "공사기간_일": int(days.group(1).replace(",", "")) if days else None,
           "추정가격_원": amt("추정가격"), "기초금액_원": amt("기초금액"), "관급자관급액_원": amt("관급자관급액"),
           "총공사금액_원": amt("총공사금액"), "신뢰도": "high",
           "근거문구": [{"항목": "연면적_m2", "원문": area.group(0) if area else ""}]}
    out = extract_llm.normalize_doc(doc)
    out.update({"_model": "mock", "_usage_in": 1000, "_usage_out": 300})
    return out


def run():
    server = MockG2B()
    base = server.start()
    os.environ["NO_PROXY"] = "127.0.0.1,localhost"
    tmp = tempfile.mkdtemp(prefix="g2b_e2e_")
    cfg = yaml.safe_load(open(os.path.join(ROOT, "config.yaml"), encoding="utf-8"))
    cfg["api"]["base_url"] = base + "/BidPublicInfoService"
    cfg["api"]["sleep_between_calls_sec"] = 0
    cfg["api"]["max_retries"] = 2
    cfg["api"]["use_license_limit"] = True                          # 가짜 서버로 면허제한 경로도 검증
    cfg["api"]["use_awards"] = True                                 # 낙찰정보서비스(별도 End Point) 경로도 검증
    cfg["api"]["award_base_url"] = base + "/ScsbidInfoService"
    cfg["period"] = {"start": "2024-01", "end": "2024-04"}          # 2024-04 는 데이터 없음(03) 경로
    cfg["paths"] = {"data_dir": "data", "raw_dir": "data/raw", "cache_db": "data/cache.sqlite", "text_dir": "data/text",
                    "files_dir": "data/files", "out_dir": "output", "excel_name": "공사비DB.xlsx"}
    cfg_path = os.path.join(tmp, "config.yaml")
    with open(cfg_path, "w", encoding="utf-8") as f:
        yaml.safe_dump(cfg, f, allow_unicode=True)
    data = lambda *p: os.path.join(tmp, "data", *p)
    try:
        # 0) 잘못된 키 → probe 가 30 오류를 한국어로 안내하고 죽지 않음, 재시도 없음
        os.environ["G2B_SERVICE_KEY"] = "wrong-key"
        out = _run(["probe", "--ym", "2024-01", "--config", cfg_path])
        assert "API 오류 30" in out and "G2B_SERVICE_KEY" in out, out
        assert server.calls["getBidPblancListInfoCnstwk"] == 1, "키 오류는 1회만 호출"
        # 1) probe — 매핑 OK + 면허제한 필드 후보 안내
        os.environ["G2B_SERVICE_KEY"] = KEY
        out = _run(["probe", "--ym", "2024-01", "--config", cfg_path])
        assert out.count("[매핑 OK]") == 2 and "후보 lcnsLmtNm 발견" in out, out
        assert "[award_amt] OK" in out and "[award_rate] OK" in out and "[낙찰정보] 기간 조회 파라미터 {'inqryDiv': '1'}" in out, out
        # 2) collect — 4개월(1개월은 NODATA), done.txt, parquet
        out = _run(["collect", "--config", cfg_path])
        notices = pd.read_parquet(data("notices_all.parquet"))
        assert len(notices) == len(server.data["notices"]), len(notices)
        assert set(open(data("raw", "notices_done.txt")).read().split()) == {"202401", "202402", "202403", "202404"}
        assert os.path.exists(data("bsis_all.parquet"))
        awards = pd.read_parquet(data("award_all.parquet"))
        assert len(awards) == len(server.data["awards"]) and awards["_award_amt"].notna().all(), awards
        assert set(open(data("raw", "award_done.txt")).read().split()) == {"202401", "202402", "202403", "202404"}
        calls_after_collect = (server.calls["getBidPblancListInfoCnstwk"], server.calls["getScsbidListSttusCnstwk"])
        _run(["collect", "--config", cfg_path])                        # 재실행: 완료된 달은 호출 없음(낙찰 목록 포함)
        assert (server.calls["getBidPblancListInfoCnstwk"], server.calls["getScsbidListSttusCnstwk"]) == calls_after_collect
        # 3) discover — 시설 3개(가상군/다른군 문화예술회관, 가상시 야외공연장), 유지보수는 '제외' 기본값, 잡음 없음
        out = _run(["discover", "--config", cfg_path])
        assert "낙찰정보 연결: 공고 4건" in out, out            # R24010001(000·001), R24010002, R24020006
        from g2b_costdb.discover import read_review_file
        rev = read_review_file(os.path.join(tmp, "output", "facility_candidates.xlsx"))
        assert len(rev) == 3, rev[["시설ID", "시설명_후보", "수요기관", "사업유형", "검수_포함여부"]]
        assert set(rev["수요기관"]) == {"가상군", "다른군", "가상시"} and (rev["검수_포함여부"] == "포함").all()
        assert list(rev[rev["수요기관"] == "가상시"]["검색어"]) == ["야외공연장"], "두 분류에 걸려도 구체적 검색어 1건"
        ids_before = dict(zip(rev["수요기관"], rev["시설ID"]))
        _run(["discover", "--config", cfg_path])                       # 재실행: 시설ID 유지
        rev2 = read_review_file(os.path.join(tmp, "output", "facility_candidates.xlsx"))
        assert dict(zip(rev2["수요기관"], rev2["시설ID"])) == ids_before
        out = _run(["discover", "--fresh", "--config", cfg_path])       # 검수 칸 초기화, 시설ID 는 유지
        rev3 = pd.read_excel(os.path.join(tmp, "output", "facility_candidates.xlsx"), sheet_name="시설후보_검수")
        assert "[--fresh]" in out and dict(zip(rev3["수요기관"], rev3["시설ID"])) == ids_before and (rev3["검수_포함여부"] == "포함").all()
        # 4) research — 면허제한 필드 자동 탐색(lcnsLmtNm), 방수공사(유지보수) 제외, 다른군은 자기 공고만
        out = _run(["research", "--config", cfg_path])
        assert "'lcnsLmtNm'" in out, out
        rs = pd.read_parquet(data("notices_research.parquet"))
        assert "R24010008" not in set(rs["공고번호"]) and set(rs[rs["시설명"] == "다른군 문화예술회관"]["공고번호"]) == {"R24020007"}
        assert rs[rs["공고번호"] == "R24010001"]["공종"].iloc[0] == "건축", "토목건축공사업 → 건축"
        assert rs[rs["공고번호"] == "R24020004"]["공고종류"].iloc[0] == "재공고", "'(취소 후 재공고)' 는 재공고"
        # 5) dedup — 변경차수 채택, 소방은 재공고가 대표, 설치공사(소액)는 대표 제외, 기초금액 없는 다른군은 대표 유지
        out = _run(["dedup", "--config", cfg_path])
        latest = pd.read_parquet(data("notices_latest.parquet"))
        rep = latest.set_index("프로젝트키")
        gid = ids_before["가상군"]
        assert rep.loc[f"{gid}-N|건축|", "공고번호"] == "R24010001" and rep.loc[f"{gid}-N|건축|", "공고차수"] == "001"
        assert rep.loc[f"{gid}-N|소방|", "공고번호"] == "R24020004"
        assert "R24030005" not in set(latest["공고번호"]) and "소액 공고 분리" in out
        assert len(latest) == 5, latest[["프로젝트키", "공고번호"]]
        # 6) attach — HWPX/DOCX/cp949 TXT 추출, HTML 응답은 실패 기록, 재실행 시 파일 재요청 없음
        out = _run(["attach", "--config", cfg_path])
        texts = json.load(open(data("texts.json"), encoding="utf-8"))
        notes = pd.read_parquet(data("notes_attach.parquet"))
        assert "연면적 15,200㎡" in texts["R24010001-001"] and "관급자관급액" in texts["R24010001-001"] and "건축면적 6,300㎡" in texts["R24010001-001"], \
            "대표(변경 차수 001)에는 변경공고문만 있어도 000 차수의 입찰공고문·현장설명서를 합쳐 처리"
        assert "정정합니다" in texts["R24010001-001"] and "다른 차수의 첨부를 합쳐 처리한 공고 1건" in out, out
        assert "추정가격 2,600,000,000원" in texts["R24010002-000"], "cp949 텍스트 복원"
        html_note = notes[notes["URL"].str.endswith("login.html")].iloc[0]
        assert html_note["다운로드"] == "N", html_note.to_dict()
        assert "객석 1,200석" in texts["R24020006-000"]
        n_file_calls = sum(v for k, v in server.calls.items() if k.startswith("file:"))
        _run(["attach", "--config", cfg_path])
        assert sum(v for k, v in server.calls.items() if k.startswith("file:")) == n_file_calls, "이미 처리한 공고는 재요청 없음"
        n_notes = len(pd.read_parquet(data("notes_attach.parquet")))
        out = _run(["attach", "--retry-failed", "--config", cfg_path])          # 실패 파일(login.html)이 있는 공고만 다시 처리, 기록은 중복되지 않음
        notes2 = pd.read_parquet(data("notes_attach.parquet"))
        assert "[--retry-failed]" in out and len(notes2) == n_notes and not notes2.duplicated(["공고번호", "URL"]).any(), (n_notes, len(notes2))
        assert sum(v for k, v in server.calls.items() if k.startswith("file:")) == n_file_calls + 1, "실패했던 파일 1개만 다시 요청"
        assert json.load(open(data("texts.json"), encoding="utf-8")).keys() == texts.keys()
        out = _run(["attach", "--report", "--config", cfg_path])
        assert "텍스트 없는 공고 2건의 원인" in out and "다운로드 실패" in out and os.path.exists(data("notes_notext.parquet")), out
        # 7) extract — 견적만 / --yes 는 가짜 추출기로
        out = _run(["extract", "--config", cfg_path])
        assert "extract --yes" in out and not os.path.exists(data("llm_docs.json"))
        os.environ["ANTHROPIC_API_KEY"] = "dummy-for-test"
        with mock.patch.object(extract_llm, "extract_with_claude", side_effect=_fake_extract):
            out = _run(["extract", "--yes", "--config", cfg_path])
        docs = json.load(open(data("llm_docs.json"), encoding="utf-8"))
        assert docs["R24010001-001"]["연면적_m2"] == 15200.0 and docs["R24010001-001"]["관급자관급액_원"] == 2100000000
        # 파일 인수인계: 1건을 실패 상태로 만들고 export → (Cowork 역할) JSON 작성 → import
        docs["R24010002-000"] = {"_error": "모의 실패"}
        with open(data("llm_docs.json"), "w", encoding="utf-8") as f:
            json.dump(docs, f, ensure_ascii=False)
        out = _run(["extract", "--export", "--limit", "5", "--config", cfg_path])
        assert os.path.exists(data("llm_in", "R24010002-000.txt")) and "내보내기 완료: 1건" in out and "(--limit 5)" in out, out
        exported = open(data("llm_in", "R24010002-000.txt"), encoding="utf-8").read()
        fake = _fake_extract(exported.split("[공고문 텍스트]")[1], {"공고명": "가상군 문화예술회관 건립 전기공사"}, cfg["llm"])
        with open(data("llm_out", "R24010002-000.json"), "w", encoding="utf-8") as f:
            json.dump({k: v for k, v in fake.items() if not k.startswith("_")}, f, ensure_ascii=False)
        out = _run(["extract", "--import", "--config", cfg_path])
        assert "JSON 1건 반영" in out and "미추출 0건" in out, out
        assert os.path.exists(data("llm_done", "R24010002-000.json")) and os.path.exists(data("llm_done", "R24010002-000.txt")) \
            and not os.path.exists(data("llm_out", "R24010002-000.json")), "반영한 파일은 llm_done 으로 이동"
        # API(모의) 결과가 있는 R24010001-001 에 사람이 쓴 JSON 이 들어와도 덮어쓰지 않음
        os.makedirs(data("llm_out"), exist_ok=True)
        json.dump({"연면적_m2": 1, "신뢰도": "low"}, open(data("llm_out", "R24010001-001.json"), "w", encoding="utf-8"))
        out = _run(["extract", "--import", "--config", cfg_path])
        docs = json.load(open(data("llm_docs.json"), encoding="utf-8"))
        assert "이미 API 결과가 있어 유지한 공고 1건" in out and docs["R24010001-001"]["연면적_m2"] == 15200.0, out
        docs = json.load(open(data("llm_docs.json"), encoding="utf-8"))
        assert docs["R24010002-000"]["추정가격_원"] == 2600000000 and docs["R24010002-000"]["_model"] == "manual/cowork"
        # 배치: 제출 전 견적(50%)만 출력하고 제출하지 않음
        out = _run(["extract", "--batch", "--config", cfg_path])
        assert "제출할 공고가 없습니다" in out and "Batches 50% 할인" in out, out
        ver = pd.read_parquet(data("logs_verify.parquet"))
        assert ((ver["공고번호"] == "R24010001") & (ver["항목"] == "추정가격") & (ver["판정"] == "정상")).any(), ver
        assert set(ver["판정"]) <= {"정상", "경고", "오류", "참고", "미확인"}
        # 8) excel — 생성 + 수식 오류 0
        out = _run(["excel", "--config", cfg_path])
        xlsx = os.path.join(tmp, "output", "공사비DB.xlsx")
        assert os.path.exists(xlsx)
        import openpyxl
        wb = openpyxl.load_workbook(xlsx)
        assert wb["04_공사비DB_공종별"].max_row - 1 == len(latest) and wb["01_시설마스터"].max_row - 1 == 3
        ws4 = wb["04_공사비DB_공종별"]; h4 = [c.value for c in ws4[1]]
        r4 = next(r for r in range(2, ws4.max_row + 1) if ws4.cell(r, h4.index("공고번호") + 1).value == "R24010001")
        assert ws4.cell(r4, h4.index("관급자관급액_API") + 1).value == 2100000000, "API 관급자재 금액이 04 시트에"
        assert ws4.cell(r4, h4.index("예산금액_API") + 1).value == 35100000000
        assert ws4.cell(r4, h4.index("낙찰금액_API") + 1).value == 28500000000 and ws4.cell(r4, h4.index("낙찰률_API") + 1).value == 86.36, "낙찰 참고 컬럼(차수 001 로 연결)"
        assert ws4.cell(r4, h4.index("낙찰자") + 1).value == "가상건설(주)"
        assert str(ws4.cell(r4, h4.index("낙찰률(수식)") + 1).value).startswith("=IF(")
        hist_df = pd.read_parquet(data("notices_hist.parquet"))
        e_row = hist_df[hist_df["공고번호"] == "R24010002"].iloc[0]
        assert e_row["낙찰자"] == "가상전기(주)" and e_row["낙찰금액_API"] == 2574000000 and e_row["참가업체수"] == 7, "재개찰 2행 중 최신 개찰일시 행"
        assert hist_df[hist_df["공고번호"] == "R24030005"]["낙찰금액_API"].isna().all(), "낙찰 없는 공고는 공란"
        assert hist_df[hist_df["공고번호"] == "R24020004"]["이전공고번호"].iloc[0] == "R24010003" and \
            set(hist_df[hist_df["공고번호"] == "R24010001"]["부공종명"]) <= {"토목공사업 / 조경공사업", "토목공사업"}
        assert wb["06_검증로그"].max_row > 1 and wb["07_추출노트"].max_row > 1
        from tests.check_excel import check
        rc = check(xlsx)
        assert rc in (0, 2), "Excel 수식 오류"          # 2 = 검증 도구 없음(LibreOffice/formulas) → 건너뜀
        # 05 시설합산 값 검증(formulas 패키지가 있을 때): 연면적은 01 시트에서, 총공사비는 04 SUMIFS 로
        try:
            import formulas  # type: ignore
            import logging, warnings
            warnings.filterwarnings("ignore"); logging.disable(logging.CRITICAL)
            sol = formulas.ExcelModel().loads(xlsx).finish().calculate()
            logging.disable(logging.NOTSET)

            def val(sheet, coord):
                for k, v in sol.items():
                    if k.upper().endswith(f"{sheet.upper()}'!{coord}"):
                        x = getattr(v, "value", v); x = x.tolist() if hasattr(x, "tolist") else x
                        while isinstance(x, list) and x:
                            x = x[0]
                        return x
            ws5 = wb["05_공사비DB_시설합산"]; h5 = [c.value for c in ws5[1]]
            col = lambda name: openpyxl.utils.get_column_letter(h5.index(name) + 1)
            rows = {ws5.cell(r, 1).value: r for r in range(2, ws5.max_row + 1)}
            r = rows[f"{gid}-N"]
            area, total, per = val("05_공사비DB_시설합산", f"{col('연면적_m2')}{r}"), val("05_공사비DB_시설합산", f"{col('총공사비합계(수식)')}{r}"), val("05_공사비DB_시설합산", f"{col('㎡당공사비_원(수식)')}{r}")
            assert float(area) == 15200.0, area
            assert abs(float(total) - (33000000000 + 2100000000 + 2860000000 + 583000000)) < 1, total   # 건축(기초+관급)+전기+소방
            assert abs(float(per) - float(total) / 15200) < 1, per
            r3 = rows[ids_before["다른군"] + "-N"]
            assert abs(float(val("05_공사비DB_시설합산", f"{col('건축(수식)')}{r3}")) - 25000000000 * 1.1) < 1, "기초금액 없으면 추정가격×1.1"
            rate = val("04_공사비DB_공종별", f"{openpyxl.utils.get_column_letter(h4.index('낙찰률(수식)') + 1)}{r4}")
            assert abs(float(rate) - 28500000000 / 33000000000) < 1e-6, ("낙찰률(수식) = 낙찰금액/기초금액", rate)
            assert abs(float(total) - (33000000000 + 2100000000 + 2860000000 + 583000000)) < 1, "낙찰금액은 총공사비에 쓰이지 않음"
            print("05 시설합산 값 검증 OK")
        except ImportError:
            print("(formulas 패키지 없음 → 05 값 검증 생략)")
        print(f"OK: 통합 실행 통과 (임시 폴더 {tmp}, API 호출 {sum(v for k, v in server.calls.items() if not k.startswith('file:'))}회)")
    finally:
        server.stop()
        load_config.cache_clear()


if __name__ == "__main__":
    run()
